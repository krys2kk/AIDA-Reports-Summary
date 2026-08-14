"""Budowanie tabelarycznego raportu Excel: jeden wiersz = jeden komputer."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from parser import KluczPola
from tabela import ProfilKomputera

_NAGLOWEK_FILL = PatternFill("solid", fgColor="1F4E78")
_NAGLOWEK_FONT = Font(color="FFFFFF", bold=True, size=11)
_PASEK_B = PatternFill("solid", fgColor="F2F2F2")

_SZEROKOSC_MIN = 12
_SZEROKOSC_MAKS = 60
_BRAK_WARTOSCI = "---"
# Ok. 2-3 linijki tekstu przy zawijaniu (Excel ~15pt / linię). Ogranicza to
# tylko WIDOCZNĄ wysokość wiersza - cała wartość zostaje w komórce i jest
# widoczna po jej kliknięciu lub ręcznym rozciągnięciu wiersza.
_WYSOKOSC_WIERSZA = 42


def buduj_tabele(
    output_path: str | Path,
    profile: list[ProfilKomputera],
    wybrane_klucze: list[KluczPola],
    postep=None,
) -> Path:
    """Zapisuje raport .xlsx: wiersz = komputer, kolumny = Lp., Nazwa pliku,
    Nazwa komputera, System operacyjny i wybrane pola (w podanej kolejności).

    `postep` (opcjonalny) to funkcja (numer, lacznie, opis) wywoływana przy
    przetwarzaniu kolejnych komputerów - do paska postępu w GUI.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Raport"

    naglowki = ["Lp.", "Nazwa pliku", "Nazwa komputera", "System operacyjny"] + [
        k.etykieta for k in wybrane_klucze
    ]
    for kolumna, tekst in enumerate(naglowki, start=1):
        komorka = ws.cell(row=1, column=kolumna, value=tekst)
        komorka.font = _NAGLOWEK_FONT
        komorka.fill = _NAGLOWEK_FILL
        komorka.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = _WYSOKOSC_WIERSZA

    for i, profil in enumerate(profile, start=1):
        if postep:
            postep(i, len(profile), profil.nazwa_komputera)
        wiersz = i + 1
        ws.row_dimensions[wiersz].height = _WYSOKOSC_WIERSZA
        pasek = _PASEK_B if i % 2 == 0 else None

        wartosci_wiersza = [
            i,
            profil.zrodlo.name,
            profil.nazwa_komputera or _BRAK_WARTOSCI,
            profil.system_operacyjny or _BRAK_WARTOSCI,
        ] + [profil.wartosci.get(klucz, _BRAK_WARTOSCI) for klucz in wybrane_klucze]
        for kolumna, wartosc in enumerate(wartosci_wiersza, start=1):
            komorka = ws.cell(row=wiersz, column=kolumna, value=wartosc)
            komorka.alignment = Alignment(vertical="top", wrap_text=(kolumna > 4))
            if pasek:
                komorka.fill = pasek

    ostatni_wiersz = len(profile) + 1
    ostatnia_kolumna = len(naglowki)
    if ostatni_wiersz >= 1 and ostatnia_kolumna >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ostatnia_kolumna)}{ostatni_wiersz}"
        ws.freeze_panes = "A2"

    for kolumna in range(1, ostatnia_kolumna + 1):
        litera = get_column_letter(kolumna)
        szerokosc = len(str(naglowki[kolumna - 1])) + 2
        for wiersz_i in range(2, ostatni_wiersz + 1):
            wartosc = ws.cell(row=wiersz_i, column=kolumna).value
            if wartosc:
                szerokosc = max(szerokosc, len(str(wartosc)) + 2)
        ws.column_dimensions[litera].width = min(
            max(szerokosc, _SZEROKOSC_MIN), _SZEROKOSC_MAKS
        )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
